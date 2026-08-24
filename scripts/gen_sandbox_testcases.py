# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MODULES = []

m = {"name": "命令黑名单 (bash_blacklist.py)", "color": "4472C4", "cases": []}
m["cases"].append(("TC-BL-001", "直接黑名单命令被拦截",
    "系统已配置 command_black_list 包含 rm 和 curl",
    "1. 打开对话框或 Agent 会话\n2. 向 AI 发送指令：执行 rm -rf /tmp/test\n3. 观察 AI 是否调用 Bash 工具及工具返回结果",
    "工具调用被拦截，返回错误提示包含黑名单命令已被拒绝及命令名 rm，不执行删除操作",
    "异常", "高"))
m["cases"].append(("TC-BL-002", "通配符黑名单匹配",
    "系统已配置 command_black_list 包含 git*",
    "1. 打开对话框或 Agent 会话\n2. 向 AI 发送指令：运行 git-lfs pull\n3. 观察工具是否被拦截",
    "工具调用被拦截，git-lfs 被识别为匹配 git* 规则，返回拒绝提示",
    "异常", "中"))
m["cases"].append(("TC-BL-003", "命令替换嵌套 $() 检测",
    "系统已配置 command_black_list 包含 curl",
    "1. 向 AI 发送指令：执行 echo $(curl http://evil.com)\n2. 观察嵌套命令是否被识别并拦截",
    "工具调用被拦截，系统提取 $() 内的 curl 命令并命中黑名单规则，返回拒绝提示",
    "异常", "高"))
m["cases"].append(("TC-BL-004", "反引号命令替换检测",
    "系统已配置 command_black_list 包含 curl",
    "1. 向 AI 发送指令：执行 echo `curl http://evil.com`\n2. 观察反引号内命令是否被识别并拦截",
    "工具调用被拦截，系统提取反引号内的 curl 命令并命中黑名单规则",
    "异常", "高"))
m["cases"].append(("TC-BL-005", "sudo 包裹的黑名单命令被检测",
    "系统已配置 command_black_list 包含 rm",
    "1. 向 AI 发送指令：执行 sudo rm -rf /tmp\n2. 观察 sudo 包裹的命令是否被解包识别",
    "工具调用被拦截，系统解包 sudo 后识别出 rm，命中黑名单规则",
    "异常", "高"))
m["cases"].append(("TC-BL-006", "完整路径命令规范化匹配",
    "系统已配置 command_black_list 包含 curl",
    "1. 向 AI 发送指令：执行 /usr/bin/curl http://evil.com\n2. 观察带完整路径的命令是否被拦截",
    "工具调用被拦截，系统提取 basename 为 curl 后命中黑名单规则",
    "异常", "中"))
m["cases"].append(("TC-BL-007", "Windows .exe 后缀命令规范化",
    "Windows 环境，系统已配置 command_black_list 包含 curl",
    "1. 向 AI 发送指令：执行 curl.exe http://evil.com\n2. 观察 .exe 后缀命令是否被识别拦截",
    "工具调用被拦截，系统剥离 .exe 后缀后识别出 curl，命中规则",
    "异常", "中"))
m["cases"].append(("TC-BL-008", "空黑名单时命令正常放行",
    "系统配置 command_black_list 为空列表 []",
    "1. 向 AI 发送指令：执行 ls -la\n2. 观察命令是否被正常执行",
    "命令正常执行，无拦截提示，返回目录列表结果",
    "正常", "中"))
m["cases"].append(("TC-BL-009", "分号分隔多命令逐个检测",
    "系统已配置 command_black_list 包含 curl",
    "1. 向 AI 发送指令：执行 ls; curl http://evil.com\n2. 观察分号后的命令是否被检测拦截",
    "工具调用被拦截，系统解析分号后的 curl 命令并命中规则，整条命令不执行",
    "异常", "高"))
m["cases"].append(("TC-BL-010", "黑名单规则大小写不敏感",
    "系统已配置 command_black_list 包含大写 CURL",
    "1. 向 AI 发送指令：执行 curl http://test.com（小写）\n2. 观察是否仍被拦截",
    "工具调用被拦截，系统将规则和命令均转小写后比较，成功命中",
    "异常", "低"))
MODULES.append(m)

m = {"name": "高风险命令拦截 (shell_risk.py)", "color": "C00000", "cases": []}
m["cases"].append(("TC-SR-001", "systemctl stop 服务命令被拦截",
    "系统运行正常，无特殊配置（高风险规则硬编码内置）",
    "1. 向 AI 发送指令：停止 nginx 服务\n2. AI 尝试调用 Bash 工具执行 systemctl stop nginx\n3. 观察工具是否被拦截",
    "工具调用被拦截，返回提示：此操作可能停止服务，出于安全考虑无法直接执行，服务不被停止",
    "异常", "高"))
m["cases"].append(("TC-SR-002", "systemctl restart 服务命令被拦截",
    "系统运行正常，无特殊配置",
    "1. 向 AI 发送指令：重启 mysql 服务\n2. AI 尝试执行 systemctl restart mysql\n3. 观察工具是否被拦截",
    "工具调用被拦截，返回安全提示，服务不被重启",
    "异常", "高"))
m["cases"].append(("TC-SR-003", "service stop 命令被拦截",
    "系统运行正常，无特殊配置",
    "1. 向 AI 发送指令：执行 service nginx stop\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回安全提示",
    "异常", "高"))
m["cases"].append(("TC-SR-004", "kill 进程命令被拦截",
    "系统运行正常，无特殊配置",
    "1. 向 AI 发送指令：强制终止进程 PID 1234\n2. AI 尝试执行 kill -9 1234\n3. 观察工具是否被拦截",
    "工具调用被拦截，返回提示：此操作会终止进程，无法直接执行",
    "异常", "高"))
m["cases"].append(("TC-SR-005", "pkill 命令被拦截",
    "系统运行正常，无特殊配置",
    "1. 向 AI 发送指令：执行 pkill -f python\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回安全提示",
    "异常", "高"))
m["cases"].append(("TC-SR-006", "docker stop 容器被拦截",
    "Docker 已安装，系统运行正常",
    "1. 向 AI 发送指令：停止容器 my-container\n2. AI 尝试执行 docker stop my-container\n3. 观察工具是否被拦截",
    "工具调用被拦截，返回提示：此操作会停止容器，无法直接执行，容器不被停止",
    "异常", "高"))
m["cases"].append(("TC-SR-007", "docker kill 容器被拦截",
    "Docker 已安装，系统运行正常",
    "1. 向 AI 发送指令：执行 docker kill my-container\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回安全提示",
    "异常", "高"))
m["cases"].append(("TC-SR-008", "iptables 防火墙命令被拦截",
    "系统运行正常，无特殊配置",
    "1. 向 AI 发送指令：添加防火墙规则屏蔽 80 端口\n2. AI 尝试执行 iptables -A INPUT -j DROP\n3. 观察工具是否被拦截",
    "工具调用被拦截，返回提示：此操作影响网络访问，无法直接执行",
    "异常", "高"))
m["cases"].append(("TC-SR-009", "kubectl delete 资源被拦截",
    "kubectl 已安装，系统运行正常",
    "1. 向 AI 发送指令：删除 pod my-pod\n2. AI 尝试执行 kubectl delete pod my-pod\n3. 观察工具是否被拦截",
    "工具调用被拦截，返回安全提示，pod 不被删除",
    "异常", "高"))
m["cases"].append(("TC-SR-010", "stop.sh 脚本被拦截",
    "工作目录存在 stop.sh 脚本",
    "1. 向 AI 发送指令：运行 stop.sh 脚本\n2. AI 尝试执行 bash stop.sh\n3. 观察工具是否被拦截",
    "工具调用被拦截，系统识别出 stop.sh 脚本名，返回安全提示",
    "异常", "高"))
m["cases"].append(("TC-SR-011", "普通命令不被高风险规则拦截",
    "系统运行正常",
    "1. 向 AI 发送指令：列出当前目录文件\n2. AI 调用 Bash 工具执行 ls -la\n3. 观察命令是否正常执行",
    "命令正常执行，返回目录列表，无安全拦截提示",
    "正常", "中"))
m["cases"].append(("TC-SR-012", "管道中含高风险命令被检测",
    "系统运行正常",
    "1. 向 AI 发送指令：执行 echo test | tee log.txt; kill -9 1234\n2. 观察管道后的高风险命令是否被检测",
    "工具调用被拦截，系统解析整条命令后检测到 kill，返回安全提示",
    "异常", "高"))
MODULES.append(m)

m = {"name": "路径安全 (sandbox/paths.py)", "color": "70AD47", "cases": []}
m["cases"].append(("TC-PATH-001", "正常工作区路径访问成功",
    "沙箱已启动，工作区为 /workspace，文件 /workspace/data/file.txt 存在",
    "1. 向 AI 发送指令：读取 /workspace/data/file.txt 的内容\n2. AI 调用 Read 工具访问该路径\n3. 观察文件是否被正常读取",
    "文件内容正常返回，无路径安全错误",
    "正常", "高"))
m["cases"].append(("TC-PATH-002", "绝对路径 .. 穿越被拒绝",
    "沙箱已启动，工作区为 /workspace",
    "1. 向 AI 发送指令：读取 /workspace/../etc/passwd\n2. AI 调用 Read 工具\n3. 观察是否被拦截",
    "工具调用返回错误：路径逃逸沙箱根目录，文件不被读取",
    "异常", "高"))
m["cases"].append(("TC-PATH-003", "相对路径 .. 穿越被拒绝",
    "沙箱已启动，工作区为 /workspace",
    "1. 向 AI 发送指令：读取 ../../etc/shadow\n2. AI 调用 Read 工具\n3. 观察是否被拦截",
    "工具调用返回错误：路径逃逸沙箱根目录",
    "异常", "高"))
m["cases"].append(("TC-PATH-004", "多层 .. 路径穿越被拒绝",
    "沙箱已启动，工作区为 /workspace",
    "1. 向 AI 发送指令：读取 /workspace/a/b/../../../etc/passwd\n2. AI 调用 Read 工具\n3. 观察是否被拦截",
    "工具调用返回错误，多层 .. 组合后的逃逸路径被检测并拒绝",
    "异常", "高"))
m["cases"].append(("TC-PATH-005", "symlink 软链接攻击被拒绝",
    "沙箱工作区内 /workspace/link 是指向 /etc 的软链接",
    "1. 向 AI 发送指令：读取 /workspace/link/passwd\n2. AI 调用 Read 工具\n3. 观察软链接路径是否被检测拒绝",
    "工具调用返回错误：Symlink not allowed in sandbox path，不读取软链接目标",
    "异常", "高"))
m["cases"].append(("TC-PATH-006", "Unicode 非断行空格路径被规范化",
    "沙箱已启动，工作区存在文件名含 Unicode 空格的文件",
    "1. 向 AI 发送指令：读取含 Unicode 空格(U+00A0)的文件路径\n2. 观察路径是否被规范化处理",
    "路径中 Unicode 空格被替换为普通空格，文件正常访问",
    "正常", "低"))
m["cases"].append(("TC-PATH-007", "data: URL 媒体源被拒绝",
    "沙箱媒体解析功能启用",
    "1. 向 AI 发送包含 data:image/png;base64,abc 的媒体请求\n2. 观察系统是否拒绝 data: URL",
    "返回错误：data: URLs are not supported for media",
    "异常", "中"))
m["cases"].append(("TC-PATH-008", "HTTP URL 媒体源直接放行",
    "沙箱媒体解析功能启用",
    "1. 向 AI 发送包含 http://example.com/img.png 的媒体请求\n2. 观察 HTTP URL 是否被放行",
    "HTTP URL 不做路径解析，直接返回原始 URL",
    "正常", "低"))
MODULES.append(m)

m = {"name": "沙箱路径映射 (sandbox_paths.py)", "color": "ED7D31", "cases": []}
m["cases"].append(("TC-MAP-001", "容器路径正确映射到宿主机工作区",
    "沙箱已启动，容器工作目录 /workspace 映射到宿主机 /host/ws，文件 /host/ws/data.csv 存在",
    "1. 向 AI 发送指令：读取容器内 /workspace/data.csv\n2. 观察路径是否被正确映射并读取文件",
    "文件正常读取，路径被映射到 /host/ws/data.csv，无错误",
    "正常", "高"))
m["cases"].append(("TC-MAP-002", "容器路径逃逸工作区被拒绝",
    "沙箱已启动，容器工作目录为 /workspace",
    "1. 向 AI 发送指令：读取 /workspace/../etc/passwd\n2. 观察路径映射是否检测逃逸",
    "返回错误：Path escapes sandbox workspace，文件不被读取",
    "异常", "高"))
m["cases"].append(("TC-MAP-003", "上传目录路径以只读方式映射",
    "沙箱已启动，会话 sess1 已上传文件到 /workspace/uploads/chat/sess1/",
    "1. 向 AI 发送指令：读取 /workspace/uploads/chat/sess1/file.csv\n2. 观察路径映射结果及读写权限",
    "文件正常读取，read_only=True，对该路径执行写操作时返回只读错误",
    "正常", "高"))
m["cases"].append(("TC-MAP-004", "上传目录路径穿越被拒绝",
    "沙箱已启动，upload_mounts 已配置",
    "1. 向 AI 发送指令：读取 /workspace/uploads/chat/sess1/../../../etc/passwd\n2. 观察是否被拦截",
    "返回错误：Path escapes sandbox upload mount",
    "异常", "高"))
m["cases"].append(("TC-MAP-005", "outputs 路径映射到会话输出目录",
    "沙箱已启动，有效 session_id 存在",
    "1. 向 AI 发送指令：将结果写入 /workspace/outputs/report.md\n2. 观察写操作是否成功及实际写入位置",
    "写操作成功，文件实际写入当前会话对应的宿主机输出目录，允许读写",
    "正常", "高"))
m["cases"].append(("TC-MAP-006", "plugins 路径映射到项目插件目录",
    "沙箱已启动，project_plugins_dir 已配置",
    "1. 向 AI 发送指令：写入 /workspace/.smartclaw/plugins/my_tool.yaml\n2. 观察路径映射及写入结果",
    "写操作成功，文件实际写入 project_plugins_dir 对应路径",
    "正常", "高"))
m["cases"].append(("TC-MAP-007", "plugins 路径穿越被拒绝",
    "沙箱已启动，project_plugins_dir 已配置",
    "1. 向 AI 发送指令：读取 /workspace/.smartclaw/plugins/../../secret\n2. 观察是否被拦截",
    "返回错误：Path escapes project plugin directory",
    "异常", "高"))
m["cases"].append(("TC-MAP-008", "无沙箱配置时路径直接使用",
    "沙箱未启用",
    "1. 向 AI 发送指令：读取某文件\n2. 观察路径是否直接使用不做沙箱映射",
    "路径直接返回原始值，不做沙箱路径映射处理",
    "正常", "中"))
m["cases"].append(("TC-MAP-009", "跨会话 outputs 路径不可访问",
    "沙箱已启动，当前 session_id=A，宿主机存在 session B 的输出目录",
    "1. 向 AI 发送指令：读取其他会话输出目录下的文件\n2. 观察是否被拦截",
    "路径不匹配当前会话，映射失败，返回权限错误",
    "异常", "高"))
MODULES.append(m)

m = {"name": "工具策略 (tool_policy.py)", "color": "9E480E", "cases": []}
m["cases"].append(("TC-POLICY-001", "deny 列表优先于 allow 列表",
    "沙箱策略配置 allow=['bash'], deny=['bash']",
    "1. 在沙箱会话中向 AI 发送需要执行 bash 命令的指令\n2. 观察 bash 工具是否被调用",
    "bash 工具调用被拒绝，deny 优先于 allow，返回工具不可用提示",
    "异常", "高"))
m["cases"].append(("TC-POLICY-002", "allow 为空时默认允许所有工具",
    "沙箱策略配置 allow=[], deny=[]",
    "1. 在沙箱会话中分别尝试调用 bash、read、write 等工具\n2. 观察是否均可正常调用",
    "所有工具均可正常调用，无拦截",
    "正常", "高"))
m["cases"].append(("TC-POLICY-003", "工具不在 allow 列表中被拒绝",
    "沙箱策略配置 allow=['bash', 'read'], deny=[]",
    "1. 在沙箱会话中向 AI 发送需要 write 工具的指令\n2. 观察 write 工具是否被拒绝",
    "write 工具调用被拒绝，返回提示该工具在当前沙箱策略中未被授权",
    "异常", "高"))
m["cases"].append(("TC-POLICY-004", "通配符 * 在 allow 列表中允许所有工具",
    "沙箱策略配置 allow=['*'], deny=[]",
    "1. 在沙箱会话中使用任意工具\n2. 观察是否全部放行",
    "所有工具均被允许",
    "正常", "中"))
m["cases"].append(("TC-POLICY-005", "deny 通配符 * 拒绝所有工具",
    "沙箱策略配置 allow=['bash'], deny=['*']",
    "1. 在沙箱会话中向 AI 发送 bash 执行指令\n2. 观察是否被 deny * 拦截",
    "bash 工具调用被拒绝，deny * 优先于 allow 中的 bash",
    "异常", "高"))
m["cases"].append(("TC-POLICY-006", "通配符前缀模式匹配",
    "沙箱策略配置 allow=['file_*'], deny=[]",
    "1. 在沙箱会话中分别尝试调用 file_read、file_write、bash 工具\n2. 观察各工具是否按策略放行或拦截",
    "file_read 和 file_write 被允许，bash 被拒绝",
    "正常", "中"))
m["cases"].append(("TC-POLICY-007", "agent 策略覆盖 global 策略",
    "global 策略 allow=['bash']，agent 策略 allow=['read','write']",
    "1. 在该 agent 的沙箱会话中尝试调用 bash 工具\n2. 再尝试调用 read 工具\n3. 观察生效的策略",
    "bash 被拒绝，read 被允许，agent 策略覆盖 global 策略",
    "正常", "高"))
MODULES.append(m)

m = {"name": "Docker 容器构建 (docker.py)", "color": "7030A0", "cases": []}
m["cases"].append(("TC-DOCKER-001", "只读根文件系统配置生效",
    "沙箱配置 read_only_root=True，容器已启动",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：在容器根目录 / 下创建文件\n3. 观察写操作是否被拒绝",
    "写操作被 Docker 只读文件系统拒绝，返回 read-only file system 错误",
    "正常", "高"))
m["cases"].append(("TC-DOCKER-002", "网络隔离配置 none 生效",
    "沙箱配置 network='none'，容器已启动",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：执行 curl http://example.com 测试网络\n3. 观察网络请求是否失败",
    "网络请求失败，容器无网络访问能力，返回连接错误",
    "正常", "高"))
m["cases"].append(("TC-DOCKER-003", "tmpfs 挂载目录可写",
    "沙箱配置 tmpfs=['/tmp', '/var/tmp']，容器已启动",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：在 /tmp 下创建临时文件并写入内容\n3. 观察写操作是否成功",
    "文件写入 /tmp 成功（tmpfs 可写），容器重启后文件消失",
    "正常", "高"))
m["cases"].append(("TC-DOCKER-004", "内存限制生效",
    "沙箱配置 memory='256m'，容器已启动",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：运行占用超过 256MB 内存的程序\n3. 观察进程是否被 OOM 终止",
    "进程因超出内存限制被 Docker 终止，容器内返回 Killed 或 OOM 错误",
    "正常", "中"))
m["cases"].append(("TC-DOCKER-005", "进程数限制生效",
    "沙箱配置 pids_limit=10，容器已启动",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：在容器内启动大量子进程\n3. 观察超出限制后是否被拒绝",
    "进程数超出限制后，新进程无法创建，返回 fork: Resource temporarily unavailable",
    "正常", "中"))
m["cases"].append(("TC-DOCKER-006", "容器名长度不超过 63 字符",
    "配置了超长 session_key（60+ 字符）",
    "1. 使用超长 session_key 创建沙箱\n2. 查看实际创建的容器名称\n3. 使用 docker inspect 验证容器名长度",
    "容器成功创建，名称长度不超过 63 字符",
    "异常", "中"))
m["cases"].append(("TC-DOCKER-007", "上传目录 bind mount 为只读",
    "沙箱已启动，upload_mounts 已挂载",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：修改 /workspace/uploads/chat/ 下的文件内容\n3. 观察写操作是否被拒绝",
    "写操作被拒绝，返回 read-only file system 错误，上传目录以只读方式挂载",
    "正常", "高"))
m["cases"].append(("TC-DOCKER-008", "配置变更且容器空闲时触发重建",
    "容器已存在但配置 hash 变化，且超过 5 分钟未使用",
    "1. 修改沙箱配置（如更换镜像或修改内存限制）\n2. 等待超过 5 分钟\n3. 发起新的沙箱会话\n4. 观察容器是否被重建",
    "旧容器被删除，新容器按新配置重新创建",
    "异常", "高"))
m["cases"].append(("TC-DOCKER-009", "5 分钟内热容器配置变更不重建",
    "容器 5 分钟内有活跃使用记录，配置 hash 发生变化",
    "1. 修改沙箱配置\n2. 立即（5分钟内）发起新的沙箱会话\n3. 观察容器是否被重建",
    "容器不重建，继续使用现有容器，记录 config_changed_hot 日志",
    "异常", "中"))
m["cases"].append(("TC-DOCKER-010", "upload/output/plugins bind 变更强制重建热容器",
    "容器 5 分钟内有活跃使用，但 binds 中新增了 /uploads/ 或 /outputs 路径",
    "1. 在热容器状态下更新包含上传/输出目录的 bind mount 配置\n2. 发起新的沙箱会话\n3. 观察是否触发重建",
    "即使容器处于热状态，上传/输出目录 bind 变更也强制触发容器重建",
    "异常", "高"))
MODULES.append(m)

m = {"name": "会话输出隔离 (bash.py)", "color": "00B0F0", "cases": []}
m["cases"].append(("TC-OUT-001", "Bash >> 重定向写入文档文件被拦截",
    "系统正常运行，会话已建立",
    "1. 向 AI 发送指令：将分析结果保存到 report.csv\n2. AI 尝试使用 Bash >> 重定向写文件\n3. 观察 Bash 工具是否被拦截",
    "Bash 工具调用被拦截，返回错误：Generated documents must be written with the Write tool，提示使用 Write 工具替代",
    "异常", "高"))
m["cases"].append(("TC-OUT-002", "tee 写入文档文件被拦截",
    "系统正常运行，会话已建立",
    "1. 向 AI 发送指令：将输出用 tee 保存到 output.json\n2. 观察 Bash 工具是否被拦截",
    "Bash 工具调用被拦截，返回错误提示改用 Write 工具",
    "异常", "高"))
m["cases"].append(("TC-OUT-003", "写入 /tmp 的临时脚本被允许",
    "系统正常运行，会话已建立",
    "1. 向 AI 发送指令：将临时辅助脚本写到 /tmp/helper.py\n2. 观察 Bash 工具是否正常执行",
    "Bash 工具正常执行，/tmp 下的临时脚本写入不被拦截",
    "正常", "中"))
m["cases"].append(("TC-OUT-004", "脚本写入项目根目录被拦截",
    "系统正常运行，当前工作目录为项目根目录",
    "1. 向 AI 发送指令：创建一个 run.sh 脚本到项目目录\n2. AI 尝试通过 Bash 写文件\n3. 观察是否被拦截",
    "Bash 工具调用被拦截，返回错误：Temporary helper scripts must not be created in the project directory，提示使用 artifacts 目录",
    "异常", "高"))
m["cases"].append(("TC-OUT-005", "写入 ~/.smartclaw/plugins 被拦截",
    "系统正常运行，会话已建立",
    "1. 向 AI 发送指令：将插件定义写到 ~/.smartclaw/plugins/tool.yaml\n2. 观察 Bash 工具是否被拦截",
    "Bash 工具调用被拦截，返回错误：应写到项目级插件目录而非用户级目录",
    "异常", "高"))
m["cases"].append(("TC-OUT-006", "SMARTCLAW_OUTPUTS_DIR 环境变量被正确注入",
    "系统正常运行，有效 session_id 存在",
    "1. 向 AI 发送指令：在 Bash 中打印 $SMARTCLAW_OUTPUTS_DIR 环境变量\n2. 观察返回值是否为当前会话的输出目录",
    "返回当前会话对应的输出目录路径，格式为 ~/.smartclaw/workspace/outputs/YYYY-MM-DD/sess-xxx/",
    "正常", "高"))
m["cases"].append(("TC-OUT-007", "子 Agent 输出归属主会话目录",
    "主会话 main-session 启动了子 Agent 会话 sub-session",
    "1. 在子 Agent 会话中向 AI 发送写文件指令\n2. 观察文件实际写入哪个会话目录",
    "文件写入主会话 main-session 的输出目录，子 Agent 输出与主会话共享",
    "正常", "高"))
m["cases"].append(("TC-OUT-008", "嵌套输出目录文件自动迁移",
    "Python 脚本在 SMARTCLAW_OUTPUTS_DIR 下又创建了日期/session 子目录",
    "1. 向 AI 发送指令：运行一个在输出目录下错误创建子目录的 Python 脚本\n2. Bash 执行完毕后观察文件位置",
    "Bash 工具执行后触发迁移，嵌套目录中的文件被自动移回正确的输出根目录",
    "异常", "中"))
m["cases"].append(("TC-OUT-009", "输出路径显示规范化为带 session scope 路径",
    "沙箱会话已建立，Bash 输出含 /workspace/outputs/ 短路径",
    "1. 向 AI 发送指令执行一个输出 /workspace/outputs/result.csv 路径的命令\n2. 观察 Bash 工具返回给用户的路径显示",
    "显示的路径被规范化为 /workspace/outputs/YYYY-MM-DD/session-id/result.csv 形式",
    "正常", "中"))
MODULES.append(m)

m = {"name": "环境变量安全 (env_security.py)", "color": "C00000", "cases": []}
m["cases"].append(("TC-ENV-001", "LD_PRELOAD 环境变量被拒绝",
    "系统运行在宿主机执行路径（非沙箱内部）",
    "1. 向 AI 发送指令：设置 LD_PRELOAD 环境变量后执行命令\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回错误：Security Violation: LD_PRELOAD is forbidden during host execution",
    "异常", "高"))
m["cases"].append(("TC-ENV-002", "DYLD_ 前缀环境变量被拒绝",
    "系统运行在宿主机执行路径",
    "1. 向 AI 发送指令：设置 DYLD_INSERT_LIBRARIES 后执行命令\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回安全违规错误",
    "异常", "高"))
m["cases"].append(("TC-ENV-003", "PATH 修改被拒绝",
    "系统运行在宿主机执行路径",
    "1. 向 AI 发送指令：设置自定义 PATH 环境变量后执行命令\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回错误：Custom PATH variable is forbidden during host execution",
    "异常", "高"))
m["cases"].append(("TC-ENV-004", "NODE_OPTIONS 环境变量被拒绝",
    "系统运行在宿主机执行路径",
    "1. 向 AI 发送指令：设置 NODE_OPTIONS=--require /evil.js 后执行 node 命令\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回安全违规错误",
    "异常", "高"))
m["cases"].append(("TC-ENV-005", "PYTHONPATH 环境变量被拒绝",
    "系统运行在宿主机执行路径",
    "1. 向 AI 发送指令：设置 PYTHONPATH=/evil 后执行 Python 命令\n2. 观察工具是否被拦截",
    "工具调用被拦截，返回安全违规错误",
    "异常", "高"))
m["cases"].append(("TC-ENV-006", "IFS 变量被拒绝",
    "系统运行在宿主机执行路径",
    "1. 向 AI 发送指令：设置 IFS=/ 后执行命令\n2. 观察是否被拦截",
    "工具调用被拦截，IFS 变量注入攻击被阻止",
    "异常", "高"))
m["cases"].append(("TC-ENV-007", "安全的普通环境变量被允许",
    "系统运行在宿主机执行路径",
    "1. 向 AI 发送指令：设置 MY_APP_CONFIG=value 和 LANG=en_US.UTF-8 后执行命令\n2. 观察命令是否正常执行",
    "命令正常执行，普通业务环境变量不被拦截",
    "正常", "中"))
MODULES.append(m)

m = {"name": "提升执行控制 (bash.py elevated)", "color": "FF7F00", "cases": []}
m["cases"].append(("TC-ELEV-001", "未配置提升时 host=host 参数被拒绝",
    "沙箱已激活，未配置 sandbox_elevated",
    "1. 在沙箱会话中向 AI 发送需要在宿主机执行的指令\n2. AI 尝试使用 host=host 参数调用 Bash 工具\n3. 观察是否被拒绝",
    "工具调用被拒绝，返回错误：Elevated host execution is not allowed in current sandbox policy",
    "异常", "高"))
m["cases"].append(("TC-ELEV-002", "提升 enabled=False 时被拒绝",
    "沙箱已激活，sandbox_elevated.enabled=False",
    "1. 在沙箱会话中向 AI 发送需要宿主机执行的指令\n2. 观察是否被拒绝",
    "工具调用被拒绝，提升未启用",
    "异常", "高"))
m["cases"].append(("TC-ELEV-003", "提升工具列表不含 bash 时被拒绝",
    "沙箱已激活，sandbox_elevated 配置：enabled=True，tools=['read']",
    "1. 在沙箱会话中向 AI 发送 bash 宿主机执行指令\n2. 观察 bash 是否因不在提升工具列表而被拒绝",
    "工具调用被拒绝，bash 不在 tools 列表中",
    "异常", "高"))
m["cases"].append(("TC-ELEV-004", "正确配置提升后 host=host 执行成功",
    "沙箱已激活，sandbox_elevated 配置：enabled=True，tools=['bash']",
    "1. 在沙箱会话中向 AI 发送需要宿主机执行的指令\n2. AI 使用 host=host 参数调用 Bash 工具\n3. 观察命令是否在宿主机执行成功",
    "命令在宿主机成功执行，返回结果，元数据包含 elevated=True",
    "正常", "高"))
m["cases"].append(("TC-ELEV-005", "不指定 host 参数时默认使用沙箱执行",
    "沙箱已激活",
    "1. 在沙箱会话中向 AI 发送普通 bash 指令（不要求宿主机执行）\n2. 观察命令在哪里执行",
    "命令在 Docker 容器内执行，元数据包含 sandbox=True，未触发提升路径",
    "正常", "高"))
MODULES.append(m)

m = {"name": "沙箱配置解析 (config.py)", "color": "00B050", "cases": []}
m["cases"].append(("TC-CFG-001", "mode='all' 兼容转换",
    "配置文件中 sandbox.mode 设置为 all（历史配置值）",
    "1. 将配置文件中 sandbox.mode 设为 all\n2. 重启服务或重新加载配置\n3. 发起新会话观察沙箱是否被启用",
    "沙箱被正常启用，mode='all' 被解析为 'on'",
    "正常", "中"))
m["cases"].append(("TC-CFG-002", "mode='non-main' 兼容转换",
    "配置文件中 sandbox.mode 设置为 non-main（历史配置值）",
    "1. 将配置文件中 sandbox.mode 设为 non-main\n2. 重启服务或重新加载配置\n3. 发起新会话观察沙箱是否被启用",
    "沙箱被正常启用，mode='non-main' 被解析为 'on'",
    "正常", "中"))
m["cases"].append(("TC-CFG-003", "shared scope 下 agent docker 覆写被忽略",
    "全局配置 scope='shared'，某个 agent 尝试覆写 docker.image",
    "1. 配置 scope=shared，global docker.image=python:slim\n2. 在某 agent 配置中设置 docker.image=custom-image\n3. 发起该 agent 的会话，观察实际使用的镜像",
    "容器使用全局配置的 python:slim 镜像，agent 覆写被忽略",
    "正常", "高"))
m["cases"].append(("TC-CFG-004", "agent scope 下 agent docker 覆写生效",
    "scope='agent'，agent 配置有自定义 docker.image",
    "1. 配置 scope=agent，global docker.image=base-image\n2. 在 agent 配置中设置 docker.image=custom-image\n3. 发起该 agent 的会话，观察实际使用的镜像",
    "容器使用 agent 配置的 custom-image，覆写生效",
    "正常", "高"))
m["cases"].append(("TC-CFG-005", "env 合并：agent env 覆盖 global env 同名键",
    "global docker.env={A:1}，agent docker.env={A:2, B:3}",
    "1. 按上述配置设置 global 和 agent env\n2. 发起 agent 会话\n3. 在容器内打印环境变量 A 和 B",
    "A=2（agent 覆盖），B=3（agent 新增），global 的 A=1 被覆盖",
    "正常", "中"))
m["cases"].append(("TC-CFG-006", "binds 合并：global 和 agent binds 拼接",
    "global binds=['/a:/b']，agent binds=['/c:/d']",
    "1. 按上述配置设置 global 和 agent binds\n2. 发起 agent 会话\n3. 检查容器实际挂载列表",
    "容器同时挂载 /a:/b 和 /c:/d，两个 bind 均生效",
    "正常", "中"))
MODULES.append(m)

m = {"name": "上传目录安全 (uploads.py)", "color": "FFC000", "cases": []}
m["cases"].append(("TC-UP-001", "session_id 含路径分隔符 / 被拒绝",
    "系统运行正常，沙箱已启动",
    "1. 构造含路径分隔符的 session_id（如 sess/../../etc）\n2. 尝试基于该 session_id 获取上传目录\n3. 观察是否被过滤",
    "_safe_session_id 返回 None，不生成上传目录，无文件系统操作",
    "异常", "高"))
m["cases"].append(("TC-UP-002", "session_id 为 . 被拒绝",
    "系统运行正常",
    "1. 构造 session_id='.' 尝试获取上传目录\n2. 观察是否被过滤",
    "返回 None，不生成上传目录",
    "异常", "高"))
m["cases"].append(("TC-UP-003", "session_id 含反斜杠被拒绝",
    "Windows 环境，系统运行正常",
    "1. 构造含反斜杠的 session_id（如 sess\\..\\etc）\n2. 尝试获取上传目录\n3. 观察是否被过滤",
    "返回 None，不生成上传目录",
    "异常", "高"))
m["cases"].append(("TC-UP-004", "正常 session_id 生成上传挂载",
    "系统正常运行，session_id='sess-abc123'，上传目录已存在",
    "1. 使用有效 session_id 获取上传目录挂载配置\n2. 观察挂载配置是否正确生成",
    "返回 SandboxUploadMount，container_dir='/workspace/uploads/chat/sess-abc123'，read_only=True",
    "正常", "高"))
m["cases"].append(("TC-UP-005", "上传文件在容器内可读",
    "沙箱已启动，session 已上传 data.csv 文件",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：读取 /workspace/uploads/chat/{session_id}/data.csv\n3. 观察文件是否可读",
    "文件内容正常返回，上传文件在容器内可通过映射路径访问",
    "正常", "高"))
m["cases"].append(("TC-UP-006", "上传文件在容器内不可写",
    "沙箱已启动，session 已上传 data.csv 文件",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：修改上传目录下 data.csv 的内容\n3. 观察写操作是否被拒绝",
    "写操作被拒绝，返回 read-only 错误，上传目录以只读方式挂载",
    "异常", "高"))
m["cases"].append(("TC-UP-007", "提示词中宿主机路径替换为容器路径",
    "沙箱已启动，session 上传文件存在，prompt 包含宿主机路径",
    "1. 发起包含上传文件宿主机路径的对话\n2. 观察传给 AI 的提示词中路径是否被替换",
    "提示词中的宿主机路径被替换为对应的容器路径",
    "正常", "中"))
m["cases"].append(("TC-UP-008", "不存在的上传目录返回空挂载列表",
    "session 无任何上传文件，上传目录不存在",
    "1. 使用无上传文件的 session_id 获取上传挂载配置\n2. 观察返回结果",
    "返回空列表 []，不生成任何 bind mount",
    "正常", "中"))
MODULES.append(m)

m = {"name": "集成与边界测试", "color": "595959", "cases": []}
m["cases"].append(("TC-INT-001", "沙箱模式下读取 workspace 外文件被拒绝",
    "沙箱已激活，workspace_dir=/host/ws",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：读取 /etc/passwd 文件\n3. 观察是否被沙箱路径检查拒绝",
    "工具调用返回错误：Path escapes sandbox workspace，文件不被读取",
    "异常", "高"))
m["cases"].append(("TC-INT-002", "无沙箱模式下工具不做路径限制",
    "沙箱未启用",
    "1. 打开普通（无沙箱）会话\n2. 向 AI 发送指令：读取系统文件\n3. 观察是否做沙箱路径检查",
    "不做沙箱路径映射检查，访问权限由操作系统决定",
    "正常", "高"))
m["cases"].append(("TC-INT-003", "路径映射后仍尝试逃逸被二次检查阻止",
    "沙箱已激活，路径经第一次映射后仍包含逃逸意图",
    "1. 构造经过映射后仍能逃逸根目录的路径\n2. 通过工具尝试访问该路径\n3. 观察 assert_under_root 二次检查是否阻止",
    "二次路径检查阻止逃逸，返回安全错误，文件不被访问",
    "异常", "高"))
m["cases"].append(("TC-INT-004", "命令超时后容器内进程被清理",
    "沙箱已激活，容器正在运行",
    "1. 向 AI 发送指令：执行一个耗时很长的命令（如 sleep 9999），设置很短的超时\n2. 等待超时触发\n3. 检查容器内是否仍有残留进程",
    "超时后 kill_sandbox_exec 被调用，容器内进程被终止，/tmp/.smartclaw_exec_* 文件被清理",
    "异常", "高"))
m["cases"].append(("TC-INT-005", "主会话与子 Agent 输出目录共享",
    "主会话 main-session 已启动子 Agent 会话 sub-session",
    "1. 在子 Agent 会话中让 AI 生成并保存一份报告\n2. 在主会话中查看输出目录\n3. 确认报告文件位置",
    "报告文件存在于主会话 main-session 的输出目录中，子 Agent 产出与主会话共享",
    "正常", "高"))
m["cases"].append(("TC-INT-006", "沙箱环境变量在容器内正确注入",
    "沙箱已激活并启动容器",
    "1. 打开沙箱会话\n2. 向 AI 发送指令：在 Bash 中打印所有 SMARTCLAW_ 开头的环境变量\n3. 验证环境变量值",
    "容器内存在 SMARTCLAW_OUTPUTS_DIR、SMARTCLAW_WORKSPACE_DIR、SMARTCLAW_ARTIFACTS_DIR、SMARTCLAW_SESSION_ID 等变量且值正确",
    "正常", "高"))
m["cases"].append(("TC-INT-007", "高风险命令在沙箱内同样被拦截",
    "沙箱已激活，Docker 容器正在运行",
    "1. 在沙箱会话中向 AI 发送停止某个服务的指令\n2. AI 尝试在容器内执行 systemctl stop 命令\n3. 观察是否被拦截",
    "工具调用在进入沙箱执行前被高风险命令检查拦截，返回安全提示",
    "异常", "高"))
m["cases"].append(("TC-INT-008", "黑名单与高风险双重检查顺序",
    "系统配置了黑名单且高风险规则均启用",
    "1. 向 AI 发送既触发黑名单又属于高风险的命令指令\n2. 观察哪个检查先生效",
    "黑名单检查先于高风险检查执行，命令被黑名单拦截并返回黑名单提示信息",
    "正常", "中"))
m["cases"].append(("TC-INT-009", "容器 setup_command 在首次创建时执行",
    "沙箱配置了 setup_command='pip install requests'，首次创建容器",
    "1. 配置 setup_command 后首次发起沙箱会话\n2. 容器创建完成后，在沙箱内尝试 import requests\n3. 观察依赖是否已安装",
    "setup_command 在容器创建时自动执行，requests 包已安装可用",
    "正常", "中"))
m["cases"].append(("TC-INT-010", "多会话并发时输出目录相互隔离",
    "系统同时运行两个不同 session 的沙箱会话",
    "1. 在会话 A 中让 AI 生成 report_A.md\n2. 在会话 B 中让 AI 生成 report_B.md\n3. 分别检查两个会话的输出目录",
    "report_A.md 只在会话 A 的输出目录中，report_B.md 只在会话 B 的输出目录中，互不干扰",
    "正常", "高"))
MODULES.append(m)

# ── 新增：复杂集成场景 ──────────────────────────────────────────────────────────

m = {"name": "跨沙箱边界文件流转", "color": "0070C0", "cases": []}
m["cases"].append(("TC-CROSS-001", "沙箱内 Write 工具写文件，宿主机提升执行读取",
    "沙箱已激活（Docker 容器），sandbox_elevated.enabled=True，tools=['bash']",
    "1. 在沙箱会话中向 AI 发送指令：用 Write 工具将分析报告写入 /workspace/outputs/report.json\n"
    "2. 写入成功后，向 AI 发送指令：用 host=host 参数在宿主机执行 cat 命令读取该文件实际路径\n"
    "3. 观察宿主机读取结果是否与写入内容一致",
    "Write 工具写入的文件内容与宿主机提升执行读取的内容完全一致；文件宿主机路径符合 ~/.smartclaw/workspace/outputs/<date>/<session>/ 规则",
    "正常", "高"))
m["cases"].append(("TC-CROSS-002", "宿主机预置上传文件，沙箱内工具读取处理并输出结果",
    "沙箱已激活，session_id='sess-test001'，宿主机已在对应 uploads/chat/sess-test001/ 目录放置 data.csv",
    "1. 在沙箱会话中向 AI 发送指令：读取 /workspace/uploads/chat/sess-test001/data.csv 并统计行数\n"
    "2. AI 使用 Read 工具读取文件内容后用 Bash 统计行数\n"
    "3. 用 Write 工具将统计结果写入 /workspace/outputs/count.txt\n"
    "4. 宿主机验证 outputs 目录中 count.txt 内容",
    "Read 工具读取上传文件成功，Bash 统计结果正确，Write 写入 count.txt 后宿主机可在输出目录确认文件内容",
    "正常", "高"))
m["cases"].append(("TC-CROSS-003", "沙箱内 Bash 生成临时数据，沙箱内 Read 读回，Write 输出最终报告",
    "沙箱已激活，容器内 Python 可用",
    "1. 向 AI 发送指令：用 Bash 执行 Python 脚本生成 JSON 数据并写到 $SMARTCLAW_OUTPUTS_DIR/raw.json\n"
    "2. 用 Read 工具读取 /workspace/outputs/raw.json 内容\n"
    "3. 用 Write 工具将格式化后的内容写到 /workspace/outputs/final_report.md\n"
    "4. 观察每步是否成功及路径一致性",
    "Bash 写入 $SMARTCLAW_OUTPUTS_DIR/raw.json 成功；Read 读取 /workspace/outputs/raw.json 与 Bash 写入内容一致；Write 输出 final_report.md 成功；三个工具操作的实际宿主机路径指向同一会话输出目录",
    "正常", "高"))
m["cases"].append(("TC-CROSS-004", "沙箱内 Bash >> 重定向写 outputs 被拦截，改用 Write 工具成功",
    "沙箱已激活，当前在沙箱会话中",
    "1. 向 AI 发送指令：执行 echo 'result' >> /workspace/outputs/result.txt\n"
    "2. 观察 Bash 工具是否被拦截并提示改用 Write 工具\n"
    "3. 按提示改用 Write 工具写入 /workspace/outputs/result.txt\n"
    "4. 观察 Write 工具是否成功",
    "步骤 1 Bash 调用被拦截，返回错误：Generated documents must be written with the Write tool；步骤 3 Write 工具写入成功，文件在输出目录可见",
    "异常", "高"))
m["cases"].append(("TC-CROSS-005", "沙箱内写 outputs 文件，跨会话路径无法访问",
    "会话 A（sess-A）和会话 B（sess-B）均在沙箱激活状态",
    "1. 在会话 A 中用 Write 工具写入 /workspace/outputs/secret_A.txt\n"
    "2. 切换到会话 B，向 AI 发送指令：读取 /workspace/outputs/secret_A.txt\n"
    "3. 观察会话 B 是否能访问会话 A 的输出文件",
    "会话 B 的 /workspace/outputs/ 映射到会话 B 自己的输出目录，找不到 secret_A.txt；即使尝试绝对路径，沙箱路径检查也拒绝访问会话 A 的输出目录",
    "异常", "高"))
m["cases"].append(("TC-CROSS-006", "沙箱内写插件目录，宿主机提升执行验证文件存在",
    "沙箱已激活，project_plugins_dir 已配置，sandbox_elevated.enabled=True",
    "1. 在沙箱会话中用 Write 工具写入 /workspace/.smartclaw/plugins/test_tool.yaml\n"
    "2. 用 host=host 参数提升执行：在宿主机验证 project_plugins_dir/test_tool.yaml 是否存在\n"
    "3. 对比文件内容",
    "Write 写入容器路径 /workspace/.smartclaw/plugins/test_tool.yaml，通过 bind mount 实际落在宿主机 project_plugins_dir；提升执行读取文件内容与写入内容一致",
    "正常", "高"))
m["cases"].append(("TC-CROSS-007", "上传目录只读：沙箱内 Write 工具写上传目录被拒绝",
    "沙箱已激活，session 已有上传文件",
    "1. 在沙箱会话中向 AI 发送指令：用 Write 工具将内容写到 /workspace/uploads/chat/{session_id}/new.txt\n"
    "2. 观察写操作是否被拒绝",
    "Write 工具调用返回 read-only 错误，上传目录以只读方式挂载，无法写入",
    "异常", "高"))
m["cases"].append(("TC-CROSS-008", "沙箱外工具（提升）写文件，沙箱内 Read 读取",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']",
    "1. 用 host=host 参数在宿主机创建文件（写到当前会话 workspace 目录）\n"
    "2. 在沙箱内用 Read 工具通过 /workspace/ 路径读取该文件\n"
    "3. 观察内容是否一致",
    "宿主机写入文件后，沙箱内 /workspace/ 路径对应的 bind mount 下即可读取到该文件；Read 工具返回内容与写入内容一致",
    "正常", "高"))
m["cases"].append(("TC-CROSS-009", "多工具操作 artifacts 目录生命周期",
    "沙箱已激活，容器内 Python 可用",
    "1. Bash 在 $SMARTCLAW_ARTIFACTS_DIR 下生成图表文件（如 chart.png 占位符）\n"
    "2. 用 Read 工具列出 /workspace/outputs/artifacts/ 目录内容\n"
    "3. 验证 $SMARTCLAW_ARTIFACTS_DIR 与 /workspace/outputs/artifacts/ 指向同一宿主机目录",
    "Bash 写入 artifacts 目录成功；Read 在 /workspace/outputs/artifacts/ 看到文件；宿主机路径验证两者指向同一物理目录",
    "正常", "中"))
m["cases"].append(("TC-CROSS-010", "提示词路径替换后沙箱工具透明访问上传文件",
    "沙箱已激活，用户上传文件时提示词中包含宿主机路径",
    "1. 构造包含宿主机上传文件绝对路径的用户提示词（如 /home/user/.smartclaw/workspace/uploads/chat/sess-001/data.csv）\n"
    "2. 发起沙箱会话，观察 AI 收到的提示词中的路径\n"
    "3. AI 使用 Read 工具读取路径，观察是否成功",
    "提示词中宿主机路径被自动替换为 /workspace/uploads/chat/sess-001/data.csv；AI 使用该容器路径调用 Read 工具成功读取文件内容",
    "正常", "高"))
MODULES.append(m)

m = {"name": "多工具流水线协作", "color": "375623", "cases": []}
m["cases"].append(("TC-PIPE-001", "数据获取 → 处理 → 报告生成全链路",
    "沙箱已激活，容器内安装 Python 及 pandas，已上传 sales.csv",
    "1. Bash 执行 Python 脚本读取 /workspace/uploads/chat/{sess}/sales.csv，计算各品类销售额并输出 JSON\n"
    "2. Bash 脚本将 JSON 写入 $SMARTCLAW_OUTPUTS_DIR/summary.json\n"
    "3. Read 工具读取 summary.json 内容\n"
    "4. Write 工具将摘要写成 Markdown 报告到 /workspace/outputs/report.md",
    "各步骤顺次成功；最终 report.md 存在于输出目录，内容涵盖上传 CSV 数据的统计摘要；整个流程无路径错误或权限拒绝",
    "正常", "高"))
m["cases"].append(("TC-PIPE-002", "代码生成 → 写文件 → 执行 → 结果验证流水线",
    "沙箱已激活，容器内 Python 可用",
    "1. AI 用 Write 工具将生成的 Python 脚本写到 /workspace/outputs/process.py（注意：不能写到项目根目录）\n"
    "2. Bash 执行 python /workspace/outputs/process.py\n"
    "3. Bash 将执行结果写到 $SMARTCLAW_OUTPUTS_DIR/exec_result.txt\n"
    "4. Read 工具读取 exec_result.txt 验证执行结果",
    "Write 写 outputs 目录的脚本成功（不触发项目目录拦截）；Bash 执行脚本成功；exec_result.txt 内容与预期输出一致",
    "正常", "高"))
m["cases"].append(("TC-PIPE-003", "Glob 搜索 → Read 多文件读取 → Bash 聚合分析",
    "沙箱已激活，/workspace 下存在多个 .py 文件",
    "1. 用 Glob 工具搜索 /workspace/**/*.py\n"
    "2. 对返回的文件列表逐一使用 Read 工具读取\n"
    "3. Bash 执行统计脚本，汇总各文件行数，写入 $SMARTCLAW_OUTPUTS_DIR/line_counts.csv",
    "Glob 返回文件列表，路径均在 sandbox workspace 内；Read 成功读取各文件内容；Bash 聚合后 line_counts.csv 内容准确",
    "正常", "中"))
m["cases"].append(("TC-PIPE-004", "Bash 在 /tmp 写中间脚本，沙箱限制不阻止，Read 读回验证",
    "沙箱已激活，容器内 /tmp 可写（tmpfs）",
    "1. Bash 将辅助脚本写到 /tmp/helper.sh\n"
    "2. Bash 赋予执行权限并运行 /tmp/helper.sh\n"
    "3. Read 工具尝试读取 /tmp/helper.sh\n"
    "4. 观察整个流程是否通畅，不触发'项目目录禁止写临时文件'拦截",
    "步骤 1 Bash 写 /tmp 不被拦截（tmpfs 可写）；步骤 3 Read 成功读取 /tmp/helper.sh 内容；全程无沙箱路径错误",
    "正常", "中"))
m["cases"].append(("TC-PIPE-005", "工具策略限制下多工具降级处理",
    "沙箱策略 allow=['bash', 'read'], deny=[]（不允许 write）",
    "1. AI 尝试用 Write 工具保存结果 → 被策略拒绝\n"
    "2. AI 改用 Bash 的重定向写文件 → 被 bash.py 输出拦截（文档类文件不允许 Bash 写）\n"
    "3. 尝试写到 /tmp 下的临时文件 → Bash 写 /tmp 被允许\n"
    "4. Read 工具读取 /tmp 下文件 → 成功",
    "Write 工具被策略拒绝；Bash >> 写文档路径被拦截；Bash 写 /tmp 成功；Read 读 /tmp 成功；验证多层防护各自独立生效",
    "异常", "高"))
m["cases"].append(("TC-PIPE-006", "并发多步骤：Bash 后台任务 → 轮询等待 → 收集结果",
    "沙箱已激活，容器内 Python 可用",
    "1. Bash 启动后台 Python 脚本并写状态到 $SMARTCLAW_OUTPUTS_DIR/status.txt（如 running/done）\n"
    "2. Bash 轮询 status.txt 直到内容变为 done\n"
    "3. Read 工具读取 $SMARTCLAW_OUTPUTS_DIR/result.json\n"
    "4. 验证结果数据完整性",
    "后台脚本正常运行；status.txt 状态正确转换；result.json 可被 Read 工具读取；全程无路径或权限错误",
    "正常", "中"))
m["cases"].append(("TC-PIPE-007", "沙箱内批量文件操作：写多个输出 → 提升执行打包",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']",
    "1. Write 工具依次写入 /workspace/outputs/part1.txt、part2.txt、part3.txt\n"
    "2. 用 host=host 参数提升执行：在宿主机执行 tar 命令打包三个文件为 archive.tar.gz\n"
    "3. 宿主机验证 archive.tar.gz 存在且内容包含三个文件",
    "Write 写入三个文件均成功；提升执行的 tar 命令读到三个文件（路径通过 identity mount 或宿主机绝对路径访问）；archive.tar.gz 内容完整",
    "正常", "高"))
m["cases"].append(("TC-PIPE-008", "Read 读取上传 CSV → Bash Python 清洗 → Write 输出清洗结果",
    "沙箱已激活，上传目录有 dirty_data.csv（含缺失值和重复行）",
    "1. Read 工具读取 /workspace/uploads/chat/{sess}/dirty_data.csv\n"
    "2. AI 用 Write 工具将 pandas 清洗脚本写到 /workspace/outputs/clean.py\n"
    "3. Bash 执行 python /workspace/outputs/clean.py，从上传目录读 CSV，输出清洗后数据到 $SMARTCLAW_OUTPUTS_DIR/clean_data.csv\n"
    "4. Read 工具读取 clean_data.csv 验证清洗结果",
    "Read 读取原始 CSV 成功（只读上传目录）；clean.py 写入 outputs 成功；Bash 脚本在容器内访问上传文件并写到输出目录成功；最终 clean_data.csv 行数/列数符合预期",
    "正常", "高"))
m["cases"].append(("TC-PIPE-009", "Bash 环境变量注入验证后执行动态路径脚本",
    "沙箱已激活，容器内环境变量已注入",
    "1. Bash 执行 echo $SMARTCLAW_OUTPUTS_DIR 验证变量存在\n"
    "2. Bash 执行 python -c \"import os; open(os.environ['SMARTCLAW_OUTPUTS_DIR']+'/env_test.txt','w').write('ok')\"\n"
    "3. Read 工具通过 /workspace/outputs/env_test.txt 读取文件\n"
    "4. 验证文件内容为 ok",
    "$SMARTCLAW_OUTPUTS_DIR 变量正确指向会话输出目录；Python 脚本通过环境变量动态拼接路径写文件成功；Read 工具读取成功，内容为 ok",
    "正常", "高"))
m["cases"].append(("TC-PIPE-010", "任务失败回滚：Write 成功后 Bash 执行失败，输出目录仍保留已写文件",
    "沙箱已激活",
    "1. Write 工具写入 /workspace/outputs/checkpoint.json（模拟任务进度存档）\n"
    "2. Bash 执行一个预期会失败的命令（如访问不存在的命令）\n"
    "3. Read 工具读取 /workspace/outputs/checkpoint.json\n"
    "4. 验证 Bash 失败不影响已成功写入的文件",
    "Bash 执行失败返回错误；checkpoint.json 仍然存在且内容完整；Write 和 Bash 相互独立，一个工具失败不影响另一个已完成的操作",
    "正常", "中"))
MODULES.append(m)

m = {"name": "多 Agent 与子任务协作", "color": "7030A0", "cases": []}
m["cases"].append(("TC-AGENT-001", "子 Agent 写文件归属主会话输出目录",
    "主会话 main-sess 已激活，sub-agent 使用 sess_root_id=main-sess 创建",
    "1. 主会话发起子 Agent\n"
    "2. 子 Agent 用 Write 工具写 /workspace/outputs/sub_report.md\n"
    "3. 在主会话侧验证 sub_report.md 存在于 main-sess 输出目录\n"
    "4. 子 Agent 会话侧验证 $SMARTCLAW_OUTPUTS_DIR 指向主会话输出目录",
    "子 Agent 写入的文件实际落到主会话输出目录；主会话可见 sub_report.md；子 Agent 的 SMARTCLAW_OUTPUTS_DIR 值等于主会话输出目录路径",
    "正常", "高"))
m["cases"].append(("TC-AGENT-002", "子 Agent 策略不能超越主会话策略",
    "主会话沙箱策略 allow=['read','write'], deny=['bash']",
    "1. 主会话发起子 Agent\n"
    "2. 子 Agent 尝试调用 Bash 工具执行命令\n"
    "3. 观察子 Agent 是否受主会话策略约束",
    "子 Agent 调用 Bash 工具被拒绝，继承主会话的 deny=['bash'] 策略；子 Agent 的 Read/Write 工具可正常调用",
    "异常", "高"))
m["cases"].append(("TC-AGENT-003", "子 Agent 读写上传文件，主会话可见输出",
    "主会话 main-sess 已上传 input.csv，子 Agent 以 main-sess 启动",
    "1. 子 Agent 使用 Read 工具读取 /workspace/uploads/chat/main-sess/input.csv\n"
    "2. 子 Agent 对数据进行分析并用 Write 写结果到 /workspace/outputs/analysis.json\n"
    "3. 主会话侧确认 analysis.json 在主输出目录\n"
    "4. 主会话使用 Read 读取 analysis.json",
    "子 Agent 可访问主会话上传文件；分析结果写到主会话输出目录；主会话成功读取子 Agent 产出；上传目录仍为只读",
    "正常", "高"))
m["cases"].append(("TC-AGENT-004", "多层子 Agent 输出目录层层归属主会话",
    "主会话 main-sess，子 Agent sub-A，子子 Agent sub-B（都指向 main-sess 根）",
    "1. sub-B 用 Write 工具写 /workspace/outputs/deep_result.txt\n"
    "2. sub-A 用 Read 工具读 /workspace/outputs/deep_result.txt\n"
    "3. main-sess 验证 deep_result.txt 在主会话输出目录\n"
    "4. 所有层的 $SMARTCLAW_OUTPUTS_DIR 均验证",
    "所有层级 Agent 的 SMARTCLAW_OUTPUTS_DIR 指向同一主会话输出目录；deep_result.txt 任何层级均可读写；宿主机只有一份文件",
    "正常", "高"))
m["cases"].append(("TC-AGENT-005", "独立子 Agent（无 sess_root_id）输出目录与主会话隔离",
    "主会话 main-sess，独立子 Agent 使用独立 session_id=sub-independent",
    "1. 独立子 Agent 写 /workspace/outputs/isolated_result.txt\n"
    "2. 主会话尝试读取 /workspace/outputs/isolated_result.txt\n"
    "3. 观察两个输出目录是否隔离",
    "独立子 Agent 的输出写到自己的输出目录；主会话在自己的输出目录找不到 isolated_result.txt；两个会话输出互不干扰",
    "异常", "高"))
m["cases"].append(("TC-AGENT-006", "主会话 Bash 处理文件，委托子 Agent Write 输出",
    "主会话沙箱已激活，子 Agent 以相同 sess_root_id 启动",
    "1. 主会话用 Bash 在容器内处理数据并输出中间结果到 $SMARTCLAW_OUTPUTS_DIR/intermediate.json\n"
    "2. 主会话委托子 Agent 读取 intermediate.json\n"
    "3. 子 Agent 用 Read 读取并用 Write 写最终格式化报告到 /workspace/outputs/final.md\n"
    "4. 主会话用 Read 读取 final.md 验证",
    "主会话 Bash 写入中间文件；子 Agent 读取成功（同目录）；子 Agent Write 写最终报告；主会话可读取最终报告；完整协作链路无路径错误",
    "正常", "高"))
m["cases"].append(("TC-AGENT-007", "子 Agent 超时后主会话输出目录状态",
    "主会话 main-sess，子 Agent 执行长时间任务后超时",
    "1. 子 Agent 用 Write 写入 step1_done.txt（超时前完成）\n"
    "2. 子 Agent 执行 Bash 长时间任务并超时\n"
    "3. 超时后主会话检查输出目录中 step1_done.txt 是否仍存在",
    "子 Agent 超时不影响已写入的文件；step1_done.txt 在输出目录仍然完整；超时仅终止后续执行，不回滚已成功的文件操作",
    "正常", "中"))
m["cases"].append(("TC-AGENT-008", "嵌套输出子目录自动迁移到正确位置",
    "子 Agent Python 脚本错误地在 SMARTCLAW_OUTPUTS_DIR 下创建 date/session 子目录",
    "1. 子 Agent Bash 执行 Python 脚本，脚本将文件写到 $SMARTCLAW_OUTPUTS_DIR/2026-01-01/sub-sess/report.csv（嵌套子目录）\n"
    "2. 触发 Bash 工具的嵌套目录迁移逻辑\n"
    "3. 主会话 Read 读取 /workspace/outputs/report.csv（无嵌套层）\n"
    "4. 验证文件被迁移到正确位置",
    "Bash 工具执行完毕后自动迁移 report.csv 到输出根目录；/workspace/outputs/report.csv 可被 Read 访问；嵌套子目录被清理",
    "异常", "中"))
MODULES.append(m)

m = {"name": "提升执行与宿主机交互", "color": "C00000", "cases": []}
m["cases"].append(("TC-ELEV-ADV-001", "沙箱内 Bash 生成，提升 Bash 宿主机后处理",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']，宿主机安装 jq",
    "1. 沙箱内 Bash 生成 JSON 文件到 $SMARTCLAW_OUTPUTS_DIR/data.json\n"
    "2. 用 host=host 提升执行：宿主机 jq 命令处理 data.json（路径为宿主机输出目录绝对路径）\n"
    "3. 提升执行将处理结果写到同目录 data_processed.json\n"
    "4. 沙箱内 Read 工具读取 /workspace/outputs/data_processed.json",
    "沙箱 Bash 写文件成功；提升执行 jq 读取宿主机路径成功；data_processed.json 通过 bind mount 反向可见于沙箱；沙箱内 Read 读取成功",
    "正常", "高"))
m["cases"].append(("TC-ELEV-ADV-002", "提升执行写 workspace，沙箱内读取",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']",
    "1. 用 host=host 提升执行：在宿主机向 workspace 目录写入 config.yaml\n"
    "2. 沙箱内用 Read 工具读取 /workspace/config.yaml\n"
    "3. 沙箱内用 Bash 执行基于 config.yaml 的命令",
    "提升执行写入宿主机 workspace 成功；沙箱内 Read 读取相同文件成功（bind mount 双向可见）；Bash 基于配置执行成功",
    "正常", "高"))
m["cases"].append(("TC-ELEV-ADV-003", "提升执行路径安全：不能写沙箱外任意位置",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']，宿主机黑名单配置 rm",
    "1. 用 host=host 提升执行：尝试执行 rm -rf /tmp/test（黑名单命令）\n"
    "2. 观察黑名单检查是否对提升执行同样生效",
    "提升执行路径同样经过黑名单检查；rm 命令被拦截，返回黑名单错误；宿主机文件不受影响",
    "异常", "高"))
m["cases"].append(("TC-ELEV-ADV-004", "提升执行 vs 沙箱执行输出目录环境变量差异",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']",
    "1. 沙箱内 Bash（无 host=host）打印 $SMARTCLAW_OUTPUTS_DIR\n"
    "2. 提升执行 Bash（host=host）打印 $SMARTCLAW_OUTPUTS_DIR\n"
    "3. 对比两个路径",
    "沙箱内输出为容器内路径（如 /workspace/outputs/...）；提升执行输出为宿主机绝对路径（如 ~/.smartclaw/workspace/outputs/...）；两个路径通过 bind mount 指向同一物理目录",
    "正常", "高"))
m["cases"].append(("TC-ELEV-ADV-005", "沙箱内+提升执行并发写同一文件，最终内容一致",
    "沙箱已激活，sandbox_elevated.enabled=True，tools=['bash']",
    "1. 沙箱内 Bash 向 $SMARTCLAW_OUTPUTS_DIR/shared.txt 写入内容 'sandbox line'\n"
    "2. 提升执行 Bash 向宿主机输出目录的 shared.txt 追加内容 'host line'\n"
    "3. Read 工具读取 /workspace/outputs/shared.txt\n"
    "4. 验证文件包含两行内容",
    "沙箱内写入第一行成功；提升执行追加第二行成功；Read 工具读到两行内容；证明 bind mount 双向写入一致性",
    "正常", "中"))
MODULES.append(m)


def make_fill(hex_color, tint=0.0):
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    if tint > 0:
        r = int(r + (255 - r) * tint)
        g = int(g + (255 - g) * tint)
        b = int(b + (255 - b) * tint)
    rgb = f"{r:02X}{g:02X}{b:02X}"
    return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")


def make_border(style="thin"):
    side = Side(style=style)
    return Border(left=side, right=side, top=side, bottom=side)


def build_excel():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("汇总")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "沙箱模块测试用例汇总"
    c.font = Font(bold=True, size=16, color="FFFFFF")
    c.fill = make_fill("1F3864")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    for col, h in enumerate(["模块名称", "用例数", "高优先级", "中优先级", "低优先级", "Sheet"], 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = make_fill("2E75B6")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = make_border()
    ws.row_dimensions[2].height = 24

    total = 0
    for ri, mod in enumerate(MODULES, 3):
        cases = mod["cases"]
        hi = sum(1 for c in cases if c[6] == "高")
        mi = sum(1 for c in cases if c[6] == "中")
        lo = sum(1 for c in cases if c[6] == "低")
        total += len(cases)
        sn = f"M{ri - 2:02d}"
        fill = make_fill(mod["color"], tint=0.7)
        for col, val in enumerate([mod["name"], len(cases), hi, mi, lo, sn], 1):
            cell = ws.cell(row=ri, column=col, value=val)
            cell.fill = fill
            cell.border = make_border()
            cell.alignment = Alignment(
                horizontal="center" if col != 1 else "left",
                vertical="center", wrap_text=True)
            if col == 1:
                cell.font = Font(bold=True, size=10)
        ws.row_dimensions[ri].height = 28

    tr = len(MODULES) + 3
    for col in range(1, 7):
        cell = ws.cell(row=tr, column=col)
        cell.fill = make_fill("1F3864", tint=0.6)
        cell.border = make_border()
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(bold=True)
    ws.cell(row=tr, column=1).value = "合计"
    ws.cell(row=tr, column=2).value = total

    for col, w in zip("ABCDEF", [44, 10, 10, 10, 10, 12]):
        ws.column_dimensions[col].width = w

    HEADERS = ["用例ID", "用例名称", "前置条件", "操作步骤", "预期结果", "用例类型", "优先级"]
    WIDTHS = [14, 32, 34, 52, 46, 10, 10]
    PRI_COLOR = {"高": "C00000", "中": "BF8F00", "低": "375623"}
    TYPE_COLOR = {"正常": "375623", "异常": "C00000"}

    for mi, mod in enumerate(MODULES, 1):
        sname = f"M{mi:02d}"
        ws2 = wb.create_sheet(sname)
        ws2.sheet_view.showGridLines = False
        ws2.freeze_panes = "A4"

        nc = len(HEADERS)
        ws2.merge_cells(f"A1:{get_column_letter(nc)}1")
        t = ws2["A1"]
        t.value = mod["name"]
        t.font = Font(bold=True, size=13, color="FFFFFF")
        t.fill = make_fill(mod["color"])
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 32

        ws2.merge_cells(f"A2:{get_column_letter(nc)}2")
        s = ws2["A2"]
        s.value = f"共 {len(mod['cases'])} 个测试用例"
        s.font = Font(size=10, color="595959")
        s.fill = make_fill(mod["color"], tint=0.85)
        s.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[2].height = 18

        for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
            cell = ws2.cell(row=3, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = make_fill(mod["color"])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = make_border()
            ws2.column_dimensions[get_column_letter(col)].width = w
        ws2.row_dimensions[3].height = 22

        for ri, case in enumerate(mod["cases"], 4):
            cid, name, pre, steps, exp, ctype, pri = case
            even = (ri % 2 == 0)
            rfill = make_fill(mod["color"], tint=0.88 if even else 0.94)
            for col, val in enumerate([cid, name, pre, steps, exp, ctype, pri], 1):
                cell = ws2.cell(row=ri, column=col, value=val)
                cell.border = make_border()
                cell.fill = rfill
                cell.alignment = Alignment(
                    vertical="top", wrap_text=True,
                    horizontal="center" if col in (1, 6, 7) else "left")
                if col == 7 and val in PRI_COLOR:
                    cell.font = Font(bold=True, color=PRI_COLOR[val], size=10)
                elif col == 6 and val in TYPE_COLOR:
                    cell.font = Font(bold=True, color=TYPE_COLOR[val], size=10)
                elif col == 1:
                    cell.font = Font(bold=True, size=9, color="1F3864")
                else:
                    cell.font = Font(size=9)
            ws2.row_dimensions[ri].height = 90

    return wb


if __name__ == "__main__":
    wb = build_excel()
    wb.save("docs/sandbox_test_cases.xlsx")
    print("OK: docs/sandbox_test_cases.xlsx")
